from pptx import Presentation
from pptx.util import Emu
pr=Presentation('apresentacao-alupar.pptx')
bad=[]
for n,sl in enumerate(pr.slides,1):
    cards=[];texts=[]
    for sh in sl.shapes:
        if sh.left is None: continue
        b=[Emu(v).inches for v in (sh.left,sh.top,sh.width,sh.height)]
        has=sh.has_text_frame and sh.text_frame.text.strip()
        if has: texts.append((b,sh.text_frame.text[:45]))
        # cartão = retângulo grande sem texto
        if not has and b[2]>1.0 and b[3]>0.5: cards.append(b)
    for tb,txt in texts:
        for c in cards:
            # o texto começa dentro do cartão mas vaza pela base ou pela lateral
            cx,cy,cw,ch=c; x,y,w,h=tb
            if x>=cx-0.02 and x<cx+cw and y>=cy-0.02 and y<cy+ch-0.05:
                if y+h>cy+ch+0.08: bad.append(f"s{n}: vaza a base do cartao ({y+h:.2f} > {cy+ch:.2f}) «{txt}»")
                if x+w>cx+cw+0.08: bad.append(f"s{n}: vaza a lateral do cartao «{txt}»")
    # margem inferior segura
    for tb,txt in texts:
        if tb[1]+tb[3]>7.34: bad.append(f"s{n}: abaixo da margem segura ({tb[1]+tb[3]:.2f}) «{txt}»")
print(len(bad),"achados de contencao")
for b in sorted(set(bad)): print(" -",b)
