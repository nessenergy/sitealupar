from pptx import Presentation
from pptx.util import Emu
import sys
pr=Presentation('apresentacao-alupar.pptx')
SW,SH=pr.slide_width,pr.slide_height
print("slides:",len(pr.slides),"  layout:",Emu(SW).inches,"x",Emu(SH).inches)
findings=[]
def ov(a,b):
    ax,ay,aw,ah=a; bx,by,bw,bh=b
    ix=min(ax+aw,bx+bw)-max(ax,bx); iy=min(ay+ah,by+bh)-max(ay,by)
    return ix,iy
CH=0.0155  # altura por pt de fonte, aprox em polegadas por linha (14pt -> .217)
for n,sl in enumerate(pr.slides,1):
    boxes=[]
    for sh in sl.shapes:
        if sh.left is None: continue
        x,y,w,h=[Emu(v).inches for v in (sh.left,sh.top,sh.width,sh.height)]
        if x<-0.01 or y<-0.01 or x+w>Emu(SW).inches+0.01 or y+h>Emu(SH).inches+0.01:
            findings.append(f"s{n}: FORA DA PAGINA {sh.shape_type} x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            txt=sh.text_frame.text
            fs=None
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.size: fs=max(fs or 0, r.font.size.pt)
            fs=fs or 12
            # estimativa grosseira de linhas
            cpl=max(1,int(w*96/(fs*0.52)))
            lines=sum(max(1,-(-len(l)//cpl)) for l in txt.split("\n"))
            need=lines*fs*1.22/72
            if need>h+0.06:
                findings.append(f"s{n}: TRANSBORDO ~{need:.2f}in em caixa h={h:.2f}in  fs={fs}  «{txt[:55]}»")
            boxes.append(((x,y,w,h),txt[:40],fs))
    for i in range(len(boxes)):
        for j in range(i+1,len(boxes)):
            ix,iy=ov(boxes[i][0],boxes[j][0])
            if ix>0.12 and iy>0.12:
                findings.append(f"s{n}: SOBREPOSICAO texto/texto {ix:.2f}x{iy:.2f}in  «{boxes[i][1]}» vs «{boxes[j][1]}»")
print(f"\n{len(findings)} achados")
for f in findings: print(" -",f)
